"""Unattended Outlook -> validated ACT snapshot -> GitHub sync.

No CFR summary emails are sent or modified by this program.
"""
from __future__ import annotations

import argparse
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from hashlib import sha256
from io import BytesIO
import importlib.util
import json
import logging
from logging.handlers import RotatingFileHandler
import math
import os
from pathlib import Path
import re
import subprocess
import sys
from tempfile import TemporaryDirectory
import time

import openpyxl

from act_snapshot import (
    activation_updates_from_parsed, build_act_snapshot, clean_launch_year,
    normalize_source_type, numeric_activation, week_sort_key,
)
from cfr_watch_analyzer import WorkbookUpload, parse_workbooks

ROOT = Path(__file__).resolve().parent
MAILER_DIR = Path(r"C:\=Codex study==\20260417-CFR Rate Bot")
REMOTE = "https://github.com/Musashi1003/cfr-watch-board-poc.git"
RECIPIENT = "wu.james@inventec.com"
RULE_VERSION = "2026-10-20-v1"
log = logging.getLogger("act_sync")


def timestamp():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def atomic_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


@contextmanager
def single_instance(path):
    """Windows releases this lock even when the process is killed."""
    path.parent.mkdir(parents=True, exist_ok=True)
    handle = path.open("a+b")
    try:
        handle.seek(0)
        if os.name == "nt":
            import msvcrt
            try:
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            except OSError:
                yield False
                return
        else:
            import fcntl
            try:
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError:
                yield False
                return
        yield True
    finally:
        handle.close()


def git(*args, cwd=None):
    env = dict(os.environ, GIT_TERMINAL_PROMPT="0", GCM_INTERACTIVE="Never")
    result = subprocess.run(
        ["git", *args], cwd=cwd, env=env, capture_output=True, timeout=180,
        creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
    )
    if result.returncode:
        # No credential/token arguments are ever supplied to git.
        raise RuntimeError("Git " + args[0] + " failed: " + result.stderr.decode("utf-8", "replace")[-1500:])
    return result.stdout


class GitHubWorkbook:
    """An isolated checkout stages only ACT; it never touches the user's checkout."""
    def __init__(self, directory, remote=REMOTE):
        self.path = Path(directory) / "github"
        self.remote = remote
        git("clone", "--depth", "1", "--single-branch", "--branch", "main", "--no-checkout", remote, str(self.path))
        git("read-tree", "HEAD", cwd=self.path)

    def read(self, revision="HEAD"):
        return git("show", f"{revision}:ACT table.xlsx", cwd=self.path)

    def publish(self, content, week):
        (self.path / "ACT table.xlsx").write_bytes(content)
        git("add", "--", "ACT table.xlsx", cwd=self.path)
        staged = git("diff", "--cached", "--name-only", cwd=self.path).decode().splitlines()
        if staged != ["ACT table.xlsx"]:
            raise ValueError(f"Unexpected staged files: {staged}")
        git("-c", "user.name=CFR ACT Automation", "-c", f"user.email={RECIPIENT}",
            "commit", "-m", f"Update {week} ACT snapshot (automatic)", cwd=self.path)
        commit = git("rev-parse", "HEAD", cwd=self.path).decode().strip()
        # A concurrent remote update rejects this push; next scheduled run starts fresh.
        git("push", "origin", "HEAD:refs/heads/main", cwd=self.path)
        return commit

    def verify(self, expected):
        git("fetch", "--depth", "1", "origin", "main", cwd=self.path)
        actual = self.read("FETCH_HEAD")
        if sha256(actual).digest() != sha256(expected).digest():
            raise ValueError("GitHub read-back differs from the validated ACT workbook")
        return git("rev-parse", "FETCH_HEAD", cwd=self.path).decode().strip()


def snapshot_rows(content):
    """Read all cells plus ACT keys; reject duplicate keys/headers and bad values."""
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    cells, values, weeks = {}, {}, {}
    try:
        for sheet in workbook:
            cells[sheet.title] = {(c.row, c.column): c.value for row in sheet for c in row if c.value is not None}
            year = clean_launch_year(sheet.title)
            if not year:
                continue
            headers = [(c.column, str(c.value)) for c in sheet[1] if re.fullmatch(r"W\d{4}", str(c.value or ""))]
            labels = [label for _, label in headers]
            if len(labels) != len(set(labels)):
                raise ValueError(f"Duplicate week headers in {sheet.title}")
            weeks[year] = labels
            seen = set()
            for row in range(2, sheet.max_row + 1):
                model = str(sheet.cell(row, 3).value or "").strip()
                if not model:
                    continue
                source = normalize_source_type(sheet.cell(row, 2).value)
                if source not in {"Gaming NB", "PC NB"}:
                    raise ValueError(f"Unknown ACT source: {source}")
                key = source, year, model
                if key in seen:
                    raise ValueError(f"Duplicate ACT model key: {key}")
                seen.add(key)
                for column, week in headers:
                    raw = sheet.cell(row, column).value
                    value = numeric_activation(raw)
                    if raw is not None and (value is None or not math.isfinite(value) or value < 0 or value != round(value)):
                        raise ValueError(f"Invalid ACT value: {key}, {week}")
                    if value is not None:
                        values[(*key, week)] = int(value)
        return cells, values, weeks
    finally:
        workbook.close()


def previous_week(week):
    start = date.fromisocalendar(2000 + int(week[1:3]), int(week[3:]), 1)
    year, number, _ = (start - timedelta(days=7)).isocalendar()
    return f"W{year % 100:02}{number:02}"


def prepare_snapshot(original, updates, week):
    if not updates or {u["week"] for u in updates} != {week}:
        raise ValueError("ACT updates are empty or have mixed weeks")
    before_cells, before, weeks = snapshot_rows(original)
    seen = set()
    expected = {}
    for update in updates:
        key = update["source_type"], update["launch_year"], update["model"], week
        value = float(update["cumulative_activation"])
        if key in seen or not math.isfinite(value) or value < 0:
            raise ValueError(f"Duplicate or invalid generated ACT: {key}")
        seen.add(key)
        value = int(round(value))
        labels = weeks.get(key[1], [])
        if previous_week(week) not in labels:
            raise ValueError(f"Missing prior week {previous_week(week)} for launch year {key[1]}")
        if any(week_sort_key(w) > week_sort_key(week) for w in labels):
            raise ValueError(f"Refusing a stale source week {week} for launch year {key[1]}")
        if key in before and before[key] != value:
            raise ValueError(f"Revised source conflicts with stored ACT: {key}, stored={before[key]}, incoming={value}")
        prior = before.get((*key[:3], previous_week(week)))
        if prior is not None and value < prior:
            raise ValueError(f"Cumulative ACT decreased: {key}, prior={prior}, incoming={value}")
        expected[key] = value
    result = build_act_snapshot(original, updates)
    after_cells, after, _ = snapshot_rows(result["xlsx_bytes"])
    for title, cells in before_cells.items():
        if any(after_cells.get(title, {}).get(pos) != val for pos, val in cells.items()):
            raise ValueError(f"Existing workbook cells changed in {title}")
    for key, value in expected.items():
        if after.get(key) != value:
            raise ValueError(f"Generated ACT does not match saved workbook: {key}")
    years = {u["launch_year"] for u in updates}
    for key, value in before.items():
        if key[1] in years and key[3] == previous_week(week):
            current = (*key[:3], week)
            if current not in after or (current not in expected and after[current] != before.get(current, value)):
                raise ValueError(f"Missing or incorrect carry-forward: {current}")
    result["counts"] = dict(Counter(k[0] for k in after if k[1] in years and k[3] == week))
    # Keep the exact original bytes when no cells changed, avoiding meaningless commits.
    result["changed"] = before_cells != after_cells
    if not result["changed"]:
        result["xlsx_bytes"] = original
    return result


def parse_sources(paths, week):
    if len(paths) != 2 or len({p.name for p in paths}) != 2:
        raise ValueError("Exactly two distinct Gaming and PC workbooks are required")
    parsed = parse_workbooks([WorkbookUpload(path=p, filename=p.name) for p in paths])
    if parsed["missing_columns"] or any(not f["rows"] for f in parsed["files"]):
        raise ValueError(f"Incomplete source workbooks: {parsed['missing_columns']}")
    if Counter(f["source_type"] for f in parsed["files"]) != Counter({"Gaming NB": 1, "PC NB": 1}):
        raise ValueError("Both Gaming NB and PC NB sources are required")
    # Each workbook may contain history; its latest raw-data week must match the email.
    for file in parsed["files"]:
        records = [r for r in parsed["records"] if r["source_file"] == file["filename"]]
        latest = max((r["Week"] for r in records), key=week_sort_key)
        if latest != week:
            raise ValueError(f"Email/file week mismatch: {file['filename']}: {latest}, expected {week}")
        if any(not clean_launch_year(r.get("開賣年度")) for r in records if r["Week"] == week):
            raise ValueError(f"Missing launch year in {file['filename']}")
    updates = activation_updates_from_parsed(parsed)
    if {u["source_type"] for u in updates} != {"Gaming NB", "PC NB"}:
        raise ValueError("ACT could not be derived from both sources")
    return updates


def load_mailer(directory):
    spec = importlib.util.spec_from_file_location("cfr_act_source_mailer", directory / "cfr_mailer.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def discover_sources(mailer):
    namespace, outlook = mailer._get_outlook_namespace()
    items = namespace.GetDefaultFolder(6).Items
    gaming = mailer._find_distinct_week_mail_map(items, mailer.GAMING_SUBJECT_PREFIX)
    pc = mailer._find_distinct_week_mail_map(items, mailer.PC_SUBJECT_PREFIX)
    common = sorted(set(gaming) & set(pc), key=mailer.week_code_to_int)
    if not common:
        raise ValueError("No common Gaming/PC source week is available")
    return gaming, pc, common, outlook


def source_signature(mails, week):
    # Detect a revised email without saving/parsing Excel on every ten-minute poll.
    identity = [(str(m.EntryID), str(m.LastModificationTime)) for m in mails]
    return sha256(json.dumps([RULE_VERSION, week, identity]).encode()).hexdigest()


def notify(outlook, subject, body):
    mail = outlook.CreateItem(0)
    mail.To = RECIPIENT
    mail.CC = ""
    mail.BCC = ""
    mail.Subject = subject
    mail.Body = body
    mail.Send()


def flush_notifications(state, path, outlook):
    for event in state.get("notifications", []):
        if not event.get("sent_at"):
            try:
                notify(outlook, event["subject"], event["body"])
            except Exception as exc:
                state["notification_error"] = str(exc)
                atomic_json(path, state)
                log.exception("ACT notification pending; will retry on next run")
                return False
            event["sent_at"] = timestamp()
            state.pop("notification_error", None)
            atomic_json(path, state)
            log.info("Independent ACT notification sent: %s", event["subject"])
    return True


def run(args):
    directory = args.state_dir
    directory.mkdir(parents=True, exist_ok=True)
    state_path = directory / "status.json"
    state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
    outlook = None
    week = "unknown"
    started = time.monotonic()
    try:
        with TemporaryDirectory(prefix="act_", dir=directory) as temporary:
            work = Path(temporary)
            if args.files:
                week = "W" + args.week.removeprefix("WK").removeprefix("W")
                paths = args.files
                signature = sha256(b"".join(p.read_bytes() for p in paths)).hexdigest()
            else:
                mailer = load_mailer(args.mailer_dir)
                gaming, pc, common, outlook = discover_sources(mailer)
                latest = common[-1]
                week = "W" + latest
                signature = source_signature([gaming[latest], pc[latest]], week)
                if not args.dry_run:
                    flush_notifications(state, state_path, outlook)
                if not args.dry_run and not args.week and state.get("source_signature") == signature and state.get("status") == "verified":
                    log.info("%s already verified; Excel parsing and GitHub sync skipped", week)
                    return 0
            github = GitHubWorkbook(work, args.remote) if not args.baseline else None
            original = args.baseline.read_bytes() if args.baseline else github.read()
            if not args.files:
                _, _, existing_weeks = snapshot_rows(original)
                latest_stored = max((w for labels in existing_weeks.values() for w in labels), key=week_sort_key)
                candidates = [w for w in common if week_sort_key("W" + w) > week_sort_key(latest_stored)]
                target = args.week.removeprefix("WK").removeprefix("W") if args.week else (candidates[0] if candidates else latest)
                if target not in common:
                    raise ValueError(f"Both source mails are required for WK{target}")
                week = "W" + target
                signature = source_signature([gaming[target], pc[target]], week)
                paths = []
                for category, mail in (("gaming", gaming[target]), ("pc", pc[target])):
                    target_dir = work / category
                    target_dir.mkdir()
                    paths.append(mailer._save_first_excel_attachment(mail, target_dir))
            updates = parse_sources(paths, week)
            result = prepare_snapshot(original, updates, week)
            report = {
                "week": week, "counts": result["counts"], "uploaded_updates": len(updates),
                "carried_forward": result["carried_forward_count"], "changed": result["changed"],
                "source_files": [{"name": p.name, "sha256": sha256(p.read_bytes()).hexdigest()} for p in paths],
                "duration_seconds": round(time.monotonic() - started, 2),
            }
            if args.dry_run:
                atomic_json(directory / "dry_run.json", dict(report, status="dry_run", checked_at=timestamp()))
                log.info("DRY RUN passed: %s", json.dumps(report, ensure_ascii=False))
                return 0
            if github is None:
                raise ValueError("A local baseline is allowed only for dry-run")
            changed = result["changed"]
            if changed:
                github.publish(result["xlsx_bytes"], week)
            commit = github.verify(result["xlsx_bytes"])
            url = "https://github.com/Musashi1003/cfr-watch-board-poc/commit/" + commit
            failed_event_id = state.get("failure_event_id")
            recovered = any(n["id"] == failed_event_id and n.get("sent_at") for n in state.get("notifications", [])) and state.get("status") == "failed"
            state.update(report, status="verified", source_signature=signature, verified_at=timestamp(),
                         commit=commit, commit_url=url, workbook_sha256=sha256(result["xlsx_bytes"]).hexdigest(),
                         failures=0, error=None, last_success_week=week)
            event_id = "success:" + week + ":" + state["workbook_sha256"] + (":" + failed_event_id if recovered else "")
            if not any(n["id"] == event_id for n in state.get("notifications", [])):
                label = "ACT 補更新成功" if recovered else "ACT 更新成功"
                state.setdefault("notifications", []).append({
                    "id": event_id, "kind": "success", "subject": f"[{label}] {week.replace('W', 'WK', 1)}",
                    "body": f"週次：{week}\nGitHub ACT 表已讀回，整份檔案 SHA-256 核對一致。\n"
                            "筆數：" + "；".join(f"{source} {count} 筆" for source, count in result["counts"].items()) + "\n"
                            f"本次新增來源值：{result['added_count']}；沿用前週：{result['carried_forward_count']}\n"
                            f"完成時間：{state['verified_at']}\nGitHub 版本：{url}\n\n"
                            "本通知只確認 GitHub ACT 資料，未驗證 Streamlit Cloud 已載入新版本。",
                })
            atomic_json(state_path, state)
            with (directory / "history.jsonl").open("a", encoding="utf-8") as handle:
                handle.write(json.dumps({k: v for k, v in state.items() if k != "notifications"}, ensure_ascii=False) + "\n")
            log.info("GitHub read-back verified: %s %s", week, commit)
            flush_notifications(state, state_path, outlook or load_mailer(args.mailer_dir)._get_outlook_namespace()[1])
            return 0
    except Exception as exc:
        log.exception("ACT sync failed for %s", week)
        if args.dry_run:
            atomic_json(directory / "dry_run.json", dict(status="failed", week=week, error=str(exc), checked_at=timestamp()))
            return 1
        # A notification failure must not turn a verified upload into a failed upload.
        if state.get("status") == "verified" and state.get("source_signature") == locals().get("signature"):
            state["notification_error"] = str(exc)
            atomic_json(state_path, state)
            return 1
        failures = state.get("failures", 0) + 1 if state.get("week") == week else 1
        state.update(status="failed", week=week, failures=failures, error=str(exc), attempted_at=timestamp())
        if failures >= 3:
            event_id = "failure:" + week + ":" + locals().get("signature", "discovery")
            state["failure_event_id"] = event_id
            if not any(n["id"] == event_id for n in state.get("notifications", [])):
                state.setdefault("notifications", []).append({
                    "id": event_id, "kind": "failure", "subject": f"[ACT 更新失敗] {week.replace('W', 'WK', 1)}",
                    "body": f"連續 {failures} 次 ACT 更新失敗，後續排程會再重試。\n"
                            f"原因：{exc}\n最後成功週次：{state.get('last_success_week', '尚無')}\n"
                            "CFR 達標率郵件內容與寄送狀態未被修改。",
                })
        atomic_json(state_path, state)
        try:
            flush_notifications(state, state_path, outlook or load_mailer(args.mailer_dir)._get_outlook_namespace()[1])
        except Exception:
            log.exception("ACT notification pending; will retry on next run")
        return 1


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true", help="Calculate and verify only; no GitHub writes or email")
    parser.add_argument("--files", nargs=2, type=Path)
    parser.add_argument("--week", help="Optional controlled source week, e.g. 2641")
    parser.add_argument("--baseline", type=Path, help="Local ACT baseline; dry-run only")
    parser.add_argument("--mailer-dir", type=Path, default=MAILER_DIR)
    parser.add_argument("--state-dir", type=Path, default=ROOT / ".act_sync")
    parser.add_argument("--remote", default=REMOTE, help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    if (args.files and not args.week) or (args.baseline and not args.dry_run):
        parser.error("--files requires --week; --baseline requires --dry-run")
    args.state_dir.mkdir(parents=True, exist_ok=True)
    handler = RotatingFileHandler(args.state_dir / "act_sync.log", maxBytes=2_000_000, backupCount=3, encoding="utf-8")
    logging.basicConfig(level=logging.INFO, handlers=[handler], format="%(asctime)s %(levelname)s %(message)s")
    with single_instance(args.state_dir / "worker.lock") as acquired:
        if not acquired:
            log.info("Another ACT worker is running; skipped")
            return 0
        return run(args)


if __name__ == "__main__":
    raise SystemExit(main())
