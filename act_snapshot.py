"""Shared ACT calculation and workbook rules for Streamlit and the local worker."""
from __future__ import annotations

from copy import copy
from io import BytesIO
import re
import openpyxl


def week_sort_key(week: str) -> tuple[int, int, str]:
  match = re.search(r"W?(\d{2})(\d{2})", week or "")
  if not match:
    return (0, 0, week or "")
  return (int(match.group(1)), int(match.group(2)), week)


def normalize_source_type(value) -> str:
  upper_value = str(value or "").upper().replace("_", "-").replace(" ", "-")
  if "GAMING" in upper_value:
    return "Gaming NB"
  if "PC" in upper_value:
    return "PC NB"
  return str(value or "").strip() or "Uploaded"


def act_table_source_type(value) -> str:
  source_type = normalize_source_type(value)
  if source_type == "Gaming NB":
    return "GAMING-NB"
  if source_type == "PC NB":
    return "PC-NB"
  return str(value or "").strip() or "Uploaded"


def clean_launch_year(value) -> str:
  match = re.search(r"(20\d{2})", str(value or ""))
  return match.group(1) if match else ""


def launch_year_from_record(record: dict) -> str:
  for key in ("\u958b\u8ce3\u5e74\u5ea6", "LAUNCH_YEAR", "OPEN_YEAR", "SALE_YEAR"):
    launch_year = clean_launch_year(record.get(key, ""))
    if launch_year:
      return launch_year
  return ""


def act_model_from_record(record: dict) -> str:
  launch_year = launch_year_from_record(record)
  if launch_year == "2025":
    model_group = str(record.get("MODEL_GROUP", "")).strip()
    if model_group:
      return model_group
  return str(record.get("ORG_MODEL(PRODUCT_DESC)", "")).strip()


def normalized_identifier(value: str) -> str:
  return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def summary_key_for_model(model: str, summary_by_model: dict[str, dict]) -> str | None:
  if model in summary_by_model:
    return model

  model_key = normalized_identifier(model)
  if not model_key:
    return None

  prefix_matches = [
    summary_model
    for summary_model in summary_by_model
    if len(normalized_identifier(summary_model)) >= 4
    and model_key.startswith(normalized_identifier(summary_model))
  ]
  if not prefix_matches:
    return None
  return max(prefix_matches, key=lambda summary_model: len(normalized_identifier(summary_model)))


def numeric_activation(value) -> float | None:
  if value is None or value == "":
    return None
  if isinstance(value, (int, float)):
    return float(value)
  text = str(value).strip().replace(",", "")
  if not text:
    return None
  try:
    return float(text)
  except ValueError:
    return None


def activation_updates_from_parsed(parsed: dict) -> list[dict]:
  records = parsed.get("records", [])
  if not records:
    return []

  summary_by_model = parsed.get("summary_by_model", {})
  summary_by_file = parsed.get("summary_by_file", {})
  weeks_by_file: dict[str, set[str]] = {}
  for record in records:
    source_file = str(record.get("source_file", "")).strip()
    week = str(record.get("Week", "")).strip()
    if source_file and week:
      weeks_by_file.setdefault(source_file, set()).add(week)
  latest_week_by_file = {
    source_file: sorted(weeks, key=week_sort_key)[-1]
    for source_file, weeks in weeks_by_file.items()
    if weeks
  }

  scope_by_week_act_model: dict[tuple[str, str, str, str], set[tuple[str, str]]] = {}
  for record in records:
    week = str(record.get("Week", "")).strip()
    source_type = normalize_source_type(record.get("source_type", ""))
    launch_year = launch_year_from_record(record)
    act_model = act_model_from_record(record)
    raw_model = str(record.get("ORG_MODEL(PRODUCT_DESC)", "")).strip()
    source_file = str(record.get("source_file", "")).strip()
    if not week or not source_type or not launch_year or not act_model or not raw_model or not source_file:
      continue
    if latest_week_by_file.get(source_file) != week:
      continue
    scope_by_week_act_model.setdefault(
      (week, source_type, launch_year, act_model),
      set(),
    ).add((raw_model, source_file))

  updates = []
  sorted_scope_items = sorted(
    scope_by_week_act_model.items(),
    key=lambda item: (week_sort_key(item[0][0]), item[0][1:]),
  )
  for (week, source_type, launch_year, act_model), raw_models in sorted_scope_items:
    activation = 0.0
    matched_models = 0
    for raw_model, source_file in raw_models:
      file_summary = summary_by_file.get(source_file) or summary_by_model
      summary_key = summary_key_for_model(raw_model, file_summary)
      if not summary_key:
        continue
      derived_act = file_summary[summary_key].get("derived_act")
      if derived_act is None:
        continue
      activation += float(derived_act)
      matched_models += 1

    if not matched_models:
      continue

    updates.append(
      {
        "source_type": source_type,
        "launch_year": launch_year,
        "model": act_model,
        "week": week,
        "cumulative_activation": activation,
        "source": "upload",
        "source_files": sorted({source_file for _, source_file in raw_models}),
      }
    )
  return updates


def act_update_weeks(updates: list[dict]) -> list[str]:
  return sorted(
    {str(update.get("week", "")).strip() for update in updates if str(update.get("week", "")).strip()},
    key=week_sort_key,
  )


def act_update_week_label(updates: list[dict]) -> str:
  weeks = act_update_weeks(updates)
  if not weeks:
    return "N/A"
  if len(weeks) == 1:
    return weeks[0]
  return f"{weeks[0]}-{weeks[-1]}"


def worksheet_header_map(worksheet) -> dict[str, int]:
  header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
  if not header_row:
    return {}
  return {
    normalized_identifier(value): index + 1
    for index, value in enumerate(header_row)
    if str(value or "").strip()
  }


def ensure_act_sheet(workbook, launch_year: str):
  sheet_name = f"{launch_year} ACT"
  if sheet_name in workbook.sheetnames:
    return workbook[sheet_name]

  worksheet = workbook.create_sheet(sheet_name)
  model_header = "MODEL_GROUP" if launch_year == "2025" else "ORG_MODEL(PRODUCT_DESC)"
  worksheet.append(["\u958b\u8ce3\u5e74\u5ea6", "GAMING/PC", model_header])
  return worksheet


def ensure_week_column(worksheet, week: str) -> int:
  headers = worksheet_header_map(worksheet)
  week_key = normalized_identifier(week)
  if week_key in headers:
    return headers[week_key]

  header_row = next(
    worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
    (),
  )
  last_labeled_column = max(
    (
      index
      for index, value in enumerate(header_row, start=1)
      if str(value or "").strip()
    ),
    default=0,
  )
  new_column = last_labeled_column + 1
  worksheet.cell(row=1, column=new_column).value = week
  if new_column > 1:
    source_cell = worksheet.cell(row=1, column=new_column - 1)
    target_cell = worksheet.cell(row=1, column=new_column)
    if source_cell.has_style:
      target_cell._style = copy(source_cell._style)
    target_cell.number_format = source_cell.number_format
    target_cell.alignment = copy(source_cell.alignment)
  return new_column


def previous_week_column(worksheet, week: str) -> int | None:
  target_key = week_sort_key(week)
  candidates = []
  for normalized_header, column in worksheet_header_map(worksheet).items():
    if not re.fullmatch(r"W\d{4}", normalized_header):
      continue
    if week_sort_key(normalized_header) < target_key:
      candidates.append((week_sort_key(normalized_header), column))
  if not candidates:
    return None
  return max(candidates, key=lambda item: item[0])[1]


def carry_forward_act_values(worksheet, week: str, week_column: int) -> set[int]:
  previous_column = previous_week_column(worksheet, week)
  if previous_column is None:
    return set()

  carried_rows = set()
  for row_index in range(2, worksheet.max_row + 1):
    target_cell = worksheet.cell(row=row_index, column=week_column)
    if numeric_activation(target_cell.value) is not None:
      continue
    previous_value = numeric_activation(
      worksheet.cell(row=row_index, column=previous_column).value
    )
    if previous_value is None:
      continue
    target_cell.value = int(round(previous_value))
    target_cell.number_format = "#,##0"
    carried_rows.add(row_index)
  return carried_rows


def act_table_row_key(worksheet, row_index: int, source_column: int, model_column: int) -> tuple[str, str]:
  source_type = normalize_source_type(worksheet.cell(row=row_index, column=source_column).value)
  model = str(worksheet.cell(row=row_index, column=model_column).value or "").strip()
  return source_type, model


def find_or_create_act_row(worksheet, update: dict, source_column: int, model_column: int) -> int:
  target_key = (update["source_type"], update["model"])
  for row_index in range(2, worksheet.max_row + 1):
    if act_table_row_key(worksheet, row_index, source_column, model_column) == target_key:
      return row_index

  row_index = worksheet.max_row + 1
  worksheet.cell(row=row_index, column=1).value = int(update["launch_year"])
  worksheet.cell(row=row_index, column=source_column).value = act_table_source_type(update["source_type"])
  worksheet.cell(row=row_index, column=model_column).value = update["model"]
  if row_index > 2:
    for column in range(1, worksheet.max_column + 1):
      source_cell = worksheet.cell(row=row_index - 1, column=column)
      target_cell = worksheet.cell(row=row_index, column=column)
      if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
      target_cell.number_format = source_cell.number_format
      target_cell.alignment = copy(source_cell.alignment)
  return row_index


def build_act_snapshot(source_bytes: bytes | None, updates: list[dict]) -> dict:
  if source_bytes:
    workbook = openpyxl.load_workbook(BytesIO(source_bytes))
  else:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

  added_count = 0
  kept_count = 0
  carried_forward_count = 0
  try:
    worksheet_state = {}
    for launch_year, week in sorted(
      {
        (
          str(update.get("launch_year", "")).strip(),
          str(update.get("week", "")).strip(),
        )
        for update in updates
        if str(update.get("launch_year", "")).strip()
        and str(update.get("week", "")).strip()
      },
      key=lambda item: (item[0], week_sort_key(item[1])),
    ):
      worksheet = ensure_act_sheet(workbook, launch_year)
      headers = worksheet_header_map(worksheet)
      source_column = headers.get("GAMINGPC") or 2
      model_column = (
        headers.get("ORGMODELPRODUCTDESC")
        or headers.get("MODELGROUP")
        or headers.get("MODEL")
        or 3
      )
      week_column = ensure_week_column(worksheet, week)
      originally_populated_rows = {
        row_index
        for row_index in range(2, worksheet.max_row + 1)
        if numeric_activation(worksheet.cell(row=row_index, column=week_column).value)
        is not None
      }
      carried_rows = carry_forward_act_values(worksheet, week, week_column)
      worksheet_state[(launch_year, week)] = {
        "worksheet": worksheet,
        "source_column": source_column,
        "model_column": model_column,
        "week_column": week_column,
        "originally_populated_rows": originally_populated_rows,
        "carried_rows": carried_rows,
      }

    for update in updates:
      launch_year = str(update.get("launch_year", "")).strip()
      week = str(update.get("week", "")).strip()
      if not launch_year or not week:
        continue
      state = worksheet_state[(launch_year, week)]
      worksheet = state["worksheet"]
      source_column = state["source_column"]
      model_column = state["model_column"]
      week_column = state["week_column"]
      row_index = find_or_create_act_row(worksheet, update, source_column, model_column)
      target_cell = worksheet.cell(row=row_index, column=week_column)
      if row_index in state["originally_populated_rows"]:
        kept_count += 1
        continue
      target_cell.value = int(round(update["cumulative_activation"]))
      target_cell.number_format = "#,##0"
      state["carried_rows"].discard(row_index)
      added_count += 1

    carried_forward_count = sum(
      len(state["carried_rows"])
      for state in worksheet_state.values()
    )

    output = BytesIO()
    workbook.save(output)
  finally:
    workbook.close()

  return {"xlsx_bytes": output.getvalue(), "added_count": added_count, "kept_count": kept_count, "carried_forward_count": carried_forward_count}
