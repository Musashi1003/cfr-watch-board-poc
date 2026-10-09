from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import streamlit_app


def uploaded_updates(source_path: Path, week: str, launch_year: str) -> list[dict]:
  workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
  try:
    worksheet = workbook[f"{launch_year} ACT"]
    headers = [str(value or "").strip() for value in next(
      worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
    )]
    normalized_headers = {
      streamlit_app.normalized_identifier(header): index
      for index, header in enumerate(headers)
    }
    source_index = normalized_headers["GAMINGPC"]
    model_index = (
      normalized_headers.get("ORGMODELPRODUCTDESC")
      or normalized_headers.get("MODELGROUP")
      or normalized_headers["MODEL"]
    )
    week_index = normalized_headers[streamlit_app.normalized_identifier(week)]

    updates = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
      act_qty = streamlit_app.numeric_activation(row[week_index])
      if act_qty is None:
        continue
      source_type = streamlit_app.normalize_source_type(row[source_index])
      model = str(row[model_index] or "").strip()
      if not source_type or not model:
        continue
      updates.append({
        "source_type": source_type,
        "launch_year": launch_year,
        "model": model,
        "week": week.upper(),
        "cumulative_activation": act_qty,
      })
    return updates
  finally:
    workbook.close()


def main():
  parser = argparse.ArgumentParser(
    description="Complete an ACT week using uploaded values plus prior-week carry-forward."
  )
  parser.add_argument("source", type=Path, help="Downloaded ACT table containing the new week")
  parser.add_argument("target", type=Path, help="Authoritative ACT table to update")
  parser.add_argument("--week", required=True)
  parser.add_argument("--launch-year", required=True)
  args = parser.parse_args()

  updates = uploaded_updates(args.source, args.week, args.launch_year)
  if not updates:
    raise SystemExit(f"No populated {args.week.upper()} values found in {args.source}")

  with (
    patch.object(streamlit_app, "configured_act_table_path", return_value=args.target),
    patch.object(
      streamlit_app,
      "write_bytes_to_github",
      return_value=(False, "Local controlled update."),
    ),
  ):
    result = streamlit_app.update_act_table_workbook(updates)

  xlsx_bytes = result.get("xlsx_bytes")
  if not xlsx_bytes:
    raise SystemExit(f"No updated workbook was generated: {result}")
  args.target.write_bytes(xlsx_bytes)

  summary = {
    key: value
    for key, value in result.items()
    if key not in {"xlsx_bytes"}
  }
  summary["target"] = str(args.target)
  print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
  main()
