from argparse import Namespace
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

import openpyxl

import act_sync as sync


def fixture(week="W2640", existing=None):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "2026 ACT"
    sheet.append(["開賣年度", "GAMING/PC", "ORG_MODEL(PRODUCT_DESC)", week])
    sheet.append([2026, "GAMING-NB", "A", 100])
    sheet.append([2026, "PC-NB", "B", 200])
    if existing is not None:
        sheet.cell(1, 5).value = "W2641"
        sheet.cell(2, 5).value = existing
    # Regression: blank formatting far to the right must not move new weeks.
    sheet.cell(1, 200).number_format = "0"
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def update(value=150):
    return dict(source_type="Gaming NB", launch_year="2026", model="A",
                week="W2641", cumulative_activation=value)


class SnapshotTests(unittest.TestCase):
    def test_adds_week_at_labeled_end_preserves_history_and_carries_forward(self):
        result = sync.prepare_snapshot(fixture(), [update()], "W2641")
        self.assertTrue(result["changed"])
        self.assertEqual(1, result["carried_forward_count"])
        _, values, weeks = sync.snapshot_rows(result["xlsx_bytes"])
        self.assertEqual(["W2640", "W2641"], weeks["2026"])
        self.assertEqual(150, values[("Gaming NB", "2026", "A", "W2641")])
        self.assertEqual(200, values[("PC NB", "2026", "B", "W2641")])
        self.assertEqual(100, values[("Gaming NB", "2026", "A", "W2640")])

    def test_same_week_reuses_exact_bytes(self):
        first = sync.prepare_snapshot(fixture(), [update()], "W2641")["xlsx_bytes"]
        repeated = sync.prepare_snapshot(first, [update()], "W2641")
        self.assertFalse(repeated["changed"])
        self.assertEqual(first, repeated["xlsx_bytes"])

    def test_revised_source_is_rejected_before_existing_value_is_silently_kept(self):
        with self.assertRaisesRegex(ValueError, "conflicts"):
            sync.prepare_snapshot(fixture(existing=140), [update()], "W2641")

    def test_gap_decrease_and_duplicate_are_rejected(self):
        for original, updates, message in (
            (fixture("W2639"), [update()], "Missing prior"),
            (fixture(), [update(99)], "decreased"),
            (fixture(), [update(), update()], "Duplicate"),
            (fixture(), [update(float("nan"))], "invalid"),
        ):
            with self.subTest(message=message), self.assertRaisesRegex(ValueError, message):
                sync.prepare_snapshot(original, updates, "W2641")

    def test_cross_year_previous_week_uses_iso_calendar(self):
        self.assertEqual("W2653", sync.previous_week("W2701"))


class GitTests(unittest.TestCase):
    def test_push_is_only_workbook_and_fetch_detects_changed_remote(self):
        with tempfile.TemporaryDirectory() as temporary:
            base = Path(temporary)
            remote, seed = base / "remote.git", base / "seed"
            sync.git("init", "--bare", str(remote))
            sync.git("init", "-b", "main", str(seed))
            (seed / "ACT table.xlsx").write_bytes(fixture())
            (seed / "keep.txt").write_text("preserve this", encoding="utf-8")
            sync.git("add", ".", cwd=seed)
            sync.git("-c", "user.name=Test", "-c", "user.email=test@example.invalid", "commit", "-m", "seed", cwd=seed)
            sync.git("remote", "add", "origin", str(remote), cwd=seed)
            sync.git("push", "origin", "main", cwd=seed)
            worker = sync.GitHubWorkbook(base / "worker", str(remote))
            result = sync.prepare_snapshot(worker.read(), [update()], "W2641")
            commit = worker.publish(result["xlsx_bytes"], "W2641")
            self.assertEqual(commit, worker.verify(result["xlsx_bytes"]))
            self.assertEqual(b"preserve this", sync.git("show", "HEAD:keep.txt", cwd=worker.path))
            with self.assertRaisesRegex(ValueError, "read-back differs"):
                worker.verify(b"wrong expected workbook")
            # Another publisher wins; the stale worker cannot overwrite it.
            other = sync.GitHubWorkbook(base / "other", str(remote))
            other.publish(fixture(existing=160), "W2641")
            with self.assertRaises(RuntimeError):
                worker.publish(fixture(existing=170), "W2641")


class NotificationTests(unittest.TestCase):
    def test_notification_failure_preserves_verified_status_and_retries_once(self):
        with tempfile.TemporaryDirectory() as temporary:
            state = {"status": "verified", "notifications": [{"subject": "ACT", "body": "verified"}]}
            path = Path(temporary) / "status.json"
            outlook = Mock()
            outlook.CreateItem.return_value.Send.side_effect = RuntimeError("offline")
            self.assertFalse(sync.flush_notifications(state, path, outlook))
            self.assertEqual("verified", json.loads(path.read_text(encoding="utf-8"))["status"])
            outlook.CreateItem.return_value.Send.side_effect = None
            self.assertTrue(sync.flush_notifications(state, path, outlook))
            self.assertEqual(sync.RECIPIENT, outlook.CreateItem.return_value.To)
            self.assertEqual("", outlook.CreateItem.return_value.CC)
            self.assertEqual("", outlook.CreateItem.return_value.BCC)
            sent = outlook.CreateItem.call_count
            sync.flush_notifications(state, path, outlook)
            self.assertEqual(sent, outlook.CreateItem.call_count)

    def test_three_failures_send_one_separate_alert_and_keep_retrying(self):
        with tempfile.TemporaryDirectory() as temporary:
            args = Namespace(state_dir=Path(temporary), files=None, week=None, dry_run=False,
                             mailer_dir=Path("unused"))
            mailer = Mock()
            outlook = Mock()
            mailer._get_outlook_namespace.return_value = (Mock(), outlook)
            with patch.object(sync, "load_mailer", return_value=mailer), patch.object(
                sync, "discover_sources", side_effect=RuntimeError("source unavailable")
            ):
                for attempt in range(1, 5):
                    self.assertEqual(1, sync.run(args))
                    self.assertEqual(0 if attempt < 3 else 1, outlook.CreateItem.call_count)


class WorkerTests(unittest.TestCase):
    def test_dry_run_never_publishes_or_notifies(self):
        with tempfile.TemporaryDirectory() as temporary:
            base = Path(temporary)
            paths = [base / "Gaming.xlsx", base / "PC.xlsx"]
            for path in paths:
                path.write_bytes(b"source")
            baseline = base / "baseline.xlsx"
            baseline.write_bytes(fixture())
            args = Namespace(state_dir=base / "state", files=paths, week="2641", dry_run=True,
                             baseline=baseline, mailer_dir=Path("unused"), remote="unused")
            with patch.object(sync, "parse_sources", return_value=[update()]), patch.object(sync, "notify") as notify, patch.object(sync, "GitHubWorkbook") as github:
                self.assertEqual(0, sync.run(args))
                notify.assert_not_called()
                github.assert_not_called()
            self.assertFalse((args.state_dir / "status.json").exists())

    def test_verified_poll_skips_excel_and_git_but_retries_pending_notification(self):
        with tempfile.TemporaryDirectory() as temporary:
            base = Path(temporary)
            mails = [SimpleNamespace(EntryID="gaming", LastModificationTime="unchanged"),
                     SimpleNamespace(EntryID="pc", LastModificationTime="unchanged")]
            signature = sync.source_signature(mails, "W2641")
            sync.atomic_json(base / "status.json", {
                "status": "verified", "source_signature": signature,
                "notifications": [{"subject": "ACT", "body": "already verified"}],
            })
            args = Namespace(state_dir=base, files=None, week=None, dry_run=False,
                             baseline=None, mailer_dir=Path("unused"), remote="unused")
            outlook = Mock()
            with patch.object(sync, "load_mailer"), patch.object(sync, "discover_sources", return_value=(
                {"2641": mails[0]}, {"2641": mails[1]}, ["2641"], outlook
            )), patch.object(sync, "GitHubWorkbook") as github, patch.object(sync, "parse_sources") as parse:
                self.assertEqual(0, sync.run(args))
                github.assert_not_called()
                parse.assert_not_called()
                self.assertEqual(1, outlook.CreateItem.call_count)
                self.assertEqual(0, sync.run(args))
                self.assertEqual(1, outlook.CreateItem.call_count)


if __name__ == "__main__":
    unittest.main()
