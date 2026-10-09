import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import openpyxl
from openpyxl.styles import PatternFill

import streamlit_app


class ActTableWorkbookTest(unittest.TestCase):
  def make_workbook(self, path: Path, with_partial_week: bool = False):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "2026 ACT"
    worksheet.append([
      "開賣年度",
      "GAMING/PC",
      "ORG_MODEL(PRODUCT_DESC)",
      "W2639",
      "W2640" if with_partial_week else None,
    ])
    worksheet.append([2026, "GAMING-NB", "MODEL-A", 100, 140 if with_partial_week else None])
    worksheet.append([2026, "PC-NB", "MODEL-B", 200, None])
    worksheet.cell(row=1, column=200).fill = PatternFill(
      fill_type="solid",
      fgColor="FFFFFF",
    )
    workbook.save(path)
    workbook.close()

  def run_update(self, path: Path):
    updates = [
      {
        "source_type": "Gaming NB",
        "launch_year": "2026",
        "model": "MODEL-A",
        "week": "W2640",
        "cumulative_activation": 150,
      }
    ]
    with (
      patch.object(streamlit_app, "configured_act_table_path", return_value=path),
      patch.object(
        streamlit_app,
        "write_bytes_to_github",
        return_value=(False, "GitHub token is not configured."),
      ),
    ):
      return streamlit_app.update_act_table_workbook(updates)

  def workbook_values(self, result):
    workbook = openpyxl.load_workbook(BytesIO(result["xlsx_bytes"]), data_only=True)
    try:
      worksheet = workbook["2026 ACT"]
      headers = [cell.value for cell in worksheet[1]]
      week_column = headers.index("W2640") + 1
      return week_column, worksheet.cell(2, week_column).value, worksheet.cell(3, week_column).value
    finally:
      workbook.close()

  def test_new_week_updates_uploaded_models_and_carries_forward_the_rest(self):
    with TemporaryDirectory(dir=Path.cwd()) as temp_dir:
      path = Path(temp_dir) / "ACT table.xlsx"
      self.make_workbook(path)
      result = self.run_update(path)

    self.assertEqual("download", result["status"])
    self.assertEqual(1, result["added_count"])
    self.assertEqual(1, result["carried_forward_count"])
    self.assertEqual((5, 150, 200), self.workbook_values(result))

  def test_partial_week_is_completed_without_overwriting_existing_values(self):
    with TemporaryDirectory(dir=Path.cwd()) as temp_dir:
      path = Path(temp_dir) / "ACT table.xlsx"
      self.make_workbook(path, with_partial_week=True)
      result = self.run_update(path)

    self.assertEqual("download", result["status"])
    self.assertEqual(0, result["added_count"])
    self.assertEqual(1, result["kept_count"])
    self.assertEqual(1, result["carried_forward_count"])
    self.assertEqual((5, 140, 200), self.workbook_values(result))


if __name__ == "__main__":
  unittest.main()
